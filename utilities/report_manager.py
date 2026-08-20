import os
from datetime import datetime
import cv2
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
from utilities.OCR_Worker import waiting, _rows, start_worker, _set_mapper_file_path
import threading
import time
VERBOSE_LOGS = False


def log(message):
    if VERBOSE_LOGS:
        print(message)


SCRIPT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR = os.path.join(SCRIPT_DIR, "reports")
SCREENSHOTS_DIR = os.path.join(REPORTS_DIR, "screenshots")

FONT = Font(name="Segoe UI", size=11)
HEADER_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1A73E8", end_color="1A73E8", fill_type="solid")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
IMG_SCALE = 0.7

HEADERS = [
    "S.No",
    "Block",
    "Gaushala",
    "CAMERA ID",
    "CONF",
    "SCREENSHOT",
    "RECORDING",
    "OPENED"
]

COL_WIDTHS = [
    8,      # S.No
    20,     # Block
    30,     # Gaushala
    18,     # CAMERA ID
    10,     # CONF
    60,     # SCREENSHOT
    14,     # RECORDING
    14      # OPENED
]

_counter = 0
_current = None
stop_event = threading.Event()


def _ensure():
    os.makedirs(SCREENSHOTS_DIR, exist_ok=True)


def save_card_screenshot(crop_bgr):
    global _counter, _current
    _ensure()
    _counter += 1
    fname = f"card_{_counter:04d}_{time.strftime('%Y%m%d_%H%M%S')}.png"
    cv2.imwrite(os.path.join(SCREENSHOTS_DIR, fname), crop_bgr)
    _current = {"screenshot": fname,"image":crop_bgr.copy()}
def end_cycle(recording, opened):
    global _current
    if not _current:
        return
    
    _current["recording"]=recording
    _current["opened"]=opened
    waiting.put(_current)
    _current = None


def _write(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = FONT
    cell.alignment = CENTER
    cell.border = BORDER
    return cell


def save_report(timeout=None):

    # =========================================================
    # WAIT FOR OCR WORKER
    # =========================================================

    deadline = (
        None
        if timeout is None
        else time.time() + timeout
    )

    while waiting.unfinished_tasks:

        if (
            deadline is not None
            and time.time() >= deadline
        ):
            break

        time.sleep(0.1)


    # =========================================================
    # GET OCR RESULTS
    # =========================================================

    rows = list(_rows)

    if not rows:
        log("[REPORT] No rows to save")
        return

    _ensure()


    # =========================================================
    # REPORT PATH
    # =========================================================

    today = datetime.now()

    date_display = today.strftime("%d/%m/%Y")

    path = os.path.join(
        REPORTS_DIR,
        f"report_truecloud_automation_"
        f"{today.strftime('%Y-%m-%d')}.xlsx"
    )


    # =========================================================
    # CREATE WORKBOOK
    # =========================================================

    wb = Workbook()

    ws = wb.active

    ws.title = "Camera Report"


    # =========================================================
    # COUNTS
    # =========================================================

    total_count = len(rows)

    recording_count = sum(
        1
        for r in rows
        if r.get("recording") == "Yes"
    )

    opened_count = sum(
        1
        for r in rows
        if r.get("opened") == "Yes"
    )

    TOTAL_CAMERAS = 481

    not_recording_count = (
        TOTAL_CAMERAS - recording_count
    )


    # =========================================================
    # ORANGE FILL
    # =========================================================

    ORANGE_FILL = PatternFill(
        start_color="F4B183",
        end_color="F4B183",
        fill_type="solid"
    )


    # =========================================================
    # TITLE
    # =========================================================

    ws.merge_cells("A1:H1")

    title_cell = ws["A1"]

    title_cell.value = (
        f"TrueCloud Automation Report "
        f"{date_display}"
    )

    title_cell.font = Font(
        name="Segoe UI",
        bold=True,
        size=16
    )

    title_cell.fill = ORANGE_FILL

    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    title_cell.border = BORDER

    ws.row_dimensions[1].height = 35


    # =========================================================
    # HEADERS
    # =========================================================

    HEADER_ROW = 3

    for col_idx, (header, width) in enumerate(
        zip(HEADERS, COL_WIDTHS),
        1
    ):

        cell = ws.cell(
            row=HEADER_ROW,
            column=col_idx,
            value=header
        )

        cell.font = Font(
                name="Segoe UI",
                bold=True,
                size=11,
                color="000000"
            )

        cell.fill = ORANGE_FILL
        cell.alignment = CENTER

        cell.border = BORDER

        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    ws.row_dimensions[
        HEADER_ROW
    ].height = 28


    # =========================================================
    # DATA
    # =========================================================

    for i, r in enumerate(rows):

        row = i + HEADER_ROW + 1

        # -----------------------------------------------------
        # S.NO
        # -----------------------------------------------------

        _write(
            ws,
            row,
            1,
            i + 1
        )

        # -----------------------------------------------------
        # BLOCK
        # -----------------------------------------------------

        _write(
            ws,
            row,
            2,
            r.get("block", "")
        )

        # -----------------------------------------------------
        # GAUSHALA
        # -----------------------------------------------------

        _write(
            ws,
            row,
            3,
            r.get("gaushala", "")
        )

        # -----------------------------------------------------
        # CAMERA ID
        # -----------------------------------------------------
        
        _write(
            ws,
            row,
            4,
            r["id"]
        )

        # -----------------------------------------------------
        # CONFIDENCE
        # -----------------------------------------------------

        _write(
            ws,
            row,
            5,
            r["conf"]
        )

        # -----------------------------------------------------
        # SCREENSHOT
        # -----------------------------------------------------

        img_path = os.path.join(
            SCREENSHOTS_DIR,
            r["screenshot"]
        )

        if os.path.exists(img_path):

            img = Image(img_path)

            img.width = int(
                img.width * IMG_SCALE
            )

            img.height = int(
                img.height * IMG_SCALE
            )

            ws.add_image(
                img,
                f"F{row}"
            )

            ws.row_dimensions[
                row
            ].height = max(
                img.height + 8,
                50
            )

        else:

            _write(
                ws,
                row,
                6,
                "(file not found)"
            )

            ws.row_dimensions[
                row
            ].height = 50

        # -----------------------------------------------------
        # RECORDING
        # -----------------------------------------------------

        _write(
            ws,
            row,
            7,
            r["recording"]
        )

        # -----------------------------------------------------
        # OPENED
        # -----------------------------------------------------

        _write(
            ws,
            row,
            8,
            r["opened"]
        )


    # =========================================================
    # REPORT RESULT
    # =========================================================

    result_start_row = (
        HEADER_ROW
        + len(rows)
        + 3
    )


    # =========================================================
    # REPORT RESULT HEADING
    # =========================================================

    ws.merge_cells(
        start_row=result_start_row,
        start_column=1,
        end_row=result_start_row,
        end_column=8
    )

    result_title = ws.cell(
        row=result_start_row,
        column=1,
        value=f"Report Result: {date_display}"
    )

    result_title.font = Font(
        name="Segoe UI",
        bold=True,
        size=13
    )

    result_title.fill = ORANGE_FILL

    result_title.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    result_title.border = BORDER

    ws.row_dimensions[
        result_start_row
    ].height = 28


    # =========================================================
    # REPORT RESULT TABLE HEADER
    # =========================================================

    result_header_row = result_start_row + 2

    result_headers = [
        "S.No",
        "Summery",
    ]

    for col_idx, header in enumerate(
        result_headers,
        start=1
    ):

        cell = ws.cell(
            row=result_header_row,
            column=col_idx,
            value=header
        )

        cell.font = Font(
            name="Segoe UI",
            bold=True,
            size=11,
            color="000000"
        )


        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 40

        cell.fill = ORANGE_FILL

        cell.alignment = CENTER

        cell.border = BORDER


    result_values = [
        (
            1,
            f"Cameras ({total_count} / {TOTAL_CAMERAS})"
        ),
        (
            2,
            f"{opened_count} Cameras Opened"
        ),
        (
            3,
            f"{recording_count} Cameras Recording"
        ),
        (
            4,
            f"{not_recording_count} Not Recording"
        ),
    ]


    result_row = result_header_row + 1

    for serial_no, summery in result_values:


        _write(
            ws,
            result_row,
            1,
            serial_no
        )

        _write(
            ws,
            result_row,
            2,
            summery
        )

        result_row += 1
    wb.save(path)