from openpyxl import load_workbook


class CameraLookup:

    def __init__(self, excel_path):

        self.lookup = {}

        self._load(excel_path)


    def _load(self, excel_path):

      

        wb = load_workbook(
            excel_path,
            read_only=True,
            data_only=True
        )

        total_ids = 0


        # =====================================================
        # PROCESS EVERY SHEET
        # =====================================================

        for ws in wb.worksheets:
            # -------------------------------------------------
            # READ HEADERS
            # -------------------------------------------------

            headers = {}

            for col in range(
                1,
                ws.max_column + 1
            ):

                value = ws.cell(
                    1,
                    col
                ).value

                if value is None:
                    continue

                header = str(
                    value
                ).strip().lower()

                headers[header] = col


            # -------------------------------------------------
            # EXACT SOURCE COLUMNS
            # -------------------------------------------------

            camera_id_col = headers.get(
                "camera id"
            )

            block_col = headers.get(
                "vikaskhand"
            )

            gaushala_col = headers.get(
                "goashray name"
            )


            # -------------------------------------------------
            # CAMERA ID IS REQUIRED
            # -------------------------------------------------

            if camera_id_col is None:
                continue


            # -------------------------------------------------
            # BLOCK / GAUSHALA ARE ALSO REQUIRED
            # -------------------------------------------------

            if block_col is None:

            

                continue


            if gaushala_col is None:


                continue


            # =================================================
            # CURRENT LOCATION
            # =================================================

            current_block = None

            current_gaushala = None


            # =================================================
            # READ ROWS
            # =================================================

            for row in range(
                2,
                ws.max_row + 1
            ):


                # ---------------------------------------------
                # BLOCK
                # ---------------------------------------------

                block = ws.cell(
                    row,
                    block_col
                ).value

                if block not in (
                    None,
                    ""
                ):

                    current_block = str(
                        block
                    ).strip()


                # ---------------------------------------------
                # GAUSHALA
                # ---------------------------------------------

                gaushala = ws.cell(
                    row,
                    gaushala_col
                ).value

                if gaushala not in (
                    None,
                    ""
                ):

                    current_gaushala = str(
                        gaushala
                    ).strip()


                # ---------------------------------------------
                # CAMERA ID
                # ---------------------------------------------

                camera_id = ws.cell(
                    row,
                    camera_id_col
                ).value


                if camera_id in (
                    None,
                    ""
                ):

                    continue


                # ---------------------------------------------
                # NORMALIZE CAMERA ID
                # ---------------------------------------------

                camera_id = str(
                    camera_id
                ).strip()


                # ---------------------------------------------
                # SAVE LOOKUP
                # ---------------------------------------------

                self.lookup[camera_id] = {

                    "block": current_block,

                    "gaushala": current_gaushala

                }


                total_ids += 1


        # =====================================================
        # CLOSE EXCEL
        # =====================================================

        wb.close()


        # =====================================================
        # LOG RESULT
        # =====================================================

    # =========================================================
    # GET CAMERA INFORMATION
    # =========================================================

    def get(
        self,
        camera_id
    ):

        if camera_id is None:

            return None, None


        # -----------------------------------------------------
        # Normalize ID
        # -----------------------------------------------------

        camera_id = str(
            camera_id
        ).strip()


        # -----------------------------------------------------
        # Lookup
        # -----------------------------------------------------

        data = self.lookup.get(
            camera_id
        )


        # -----------------------------------------------------
        # NOT FOUND
        # -----------------------------------------------------

        if data is None:

            return None, None


        # -----------------------------------------------------
        # RETURN
        # -----------------------------------------------------

        return (

            data.get(
                "block"
            ),

            data.get(
                "gaushala"
            )

        )